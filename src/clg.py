import argparse
import json
import sys
import time
import os
import shutil
from pathlib import Path
from datetime import datetime
from docx import Document
from docx2pdf import convert

CONFIG_FILE = Path.home() / ".HireMe_config.json"

def save_config(storage_location, template_location, template_map, output_docx_dir, output_pdf_dir, swap_words=None):
    if swap_words is None:
        swap_words = ["{{COMPANY_NAME}}", "{{ROLE}}", "{{DATE}}"]

    config = {
        "storageLocation": storage_location,
        "templateLocation": template_location,
        "outputDocxLocation": output_docx_dir,
        "outputPdfLocation": output_pdf_dir,
        "templateKeywords": template_map,
        "swapWords": swap_words
    }

    with open(CONFIG_FILE, "w") as f:
        json.dump(config, f, indent=4)
    print(f"[✓] Configuration saved to {CONFIG_FILE}")


def load_config():
    if CONFIG_FILE.exists():
        with open(CONFIG_FILE, "r") as f:
            config = json.load(f)
        if "swapWords" not in config:
            config["swapWords"] = ["{{COMPANY_NAME}}", "{{ROLE}}", "{{DATE}}"]

        return config
    else:
        print("[!] Run `HireMe --configure` first to set up your environment.")
        exit(1)

def configure():
    print("Configuring HireMe...")

    if CONFIG_FILE.exists():
        try:
            with open(CONFIG_FILE, "r") as f:
                config = json.load(f)

            required_keys = [
                "storageLocation", "templateLocation",
                "outputDocxLocation", "outputPdfLocation",
                "templateKeywords", "swapWords"
            ]

            if all(k in config and config[k] for k in required_keys):
                print("[!] Existing configuration detected.")
                proceed = input("[?] Do you want to overwrite it? [Y]es/[N]o (Default=No): ").strip().lower()
                if proceed not in ("y", "yes"):
                    print("[✓] Configuration preserved. Exiting.")
                    return
        except Exception:
            print("[!] Could not read existing config. Proceeding with fresh configuration...")

    storage_location = input("[1] Generator Files Folder Path: ").strip()
    if not os.path.exists(storage_location):
        print("[!] That path does not exist.")
        return

    template_location = input("[2] Path to Folder of .docx Templates (not in [1]): ").strip()
    if not os.path.exists(template_location):
        print("[!] That path does not exist.")
        return

    templates_dir = os.path.join(storage_location, "templates")
    output_docx_dir = os.path.join(storage_location, "docxStorage")
    output_pdf_dir = os.path.join(storage_location, "outputPDFs")

    os.makedirs(templates_dir, exist_ok=True)
    os.makedirs(output_docx_dir, exist_ok=True)
    os.makedirs(output_pdf_dir, exist_ok=True)

    template_files = [f for f in os.listdir(template_location) if f.lower().endswith('.docx')]
    if not template_files:
        print("[!] No .docx files found in the template folder.")
        return

    print(f"Copying {len(template_files)} template files to {templates_dir}...")
    for file in template_files:
        src_path = os.path.join(template_location, file)
        dst_path = os.path.join(templates_dir, file)
        shutil.copyfile(src_path, dst_path)
    print("[✓] Copy complete.")

    template_map = {}
    for i, filename in enumerate(template_files, start=1):
        while True:
            prompt = f"[>] Provide keyword for template {i} ('{filename}'): "
            keyword = input(prompt).strip()
            if keyword == "":
                print("[!] A keyword is necessary. Please enter a valid keyword.")
            elif keyword in template_map:
                print("[!] Duplicate keyword. Please enter a unique one.")
            else:
                template_map[keyword] = filename
                break  

    if not template_map:
        print("[!] No keywords assigned. Configuration aborted.")
        return

    swap_words = ["{{COMPANY_NAME}}", "{{ROLE}}", "{{DATE}}"]

    print("\n[?] Do you need to add more {{PLACEHOLDERS}} ? [Y]es/[N]o (Default=No)")
    response = input("> ").strip().lower()

    while response in ("y", "yes"):
        label = input("[>] Enter a LABEL for a new placeholder {{LABEL}}: ").strip().upper()
        if label == "":
            print("[!] Empty label skipped.")
        else:
            placeholder = f"{{{{{label}}}}}"
            if placeholder in swap_words:
                print(f"[!] Placeholder {placeholder} already exists.")
            else:
                swap_words.append(placeholder)
                print(f"[✓] Added placeholder {placeholder}")

        print("\n[?] Continue adding placeholders? [Y]es/[N]o (Default=No)")
        response = input("> ").strip().lower()

    save_config(storage_location, templates_dir, template_map, output_docx_dir, output_pdf_dir, swap_words)

    print(f"\n[✓] Setup complete. You can now generate cover letters.")
    print(f"[✓] Output DOCX files will be saved to: {output_docx_dir}")
    print(f"[✓] Output PDF files will be saved to: {output_pdf_dir}")
    print(f"[✓] Placeholders saved: {', '.join(swap_words)}")



def configureLess():
    print("Configuring HireMe...")

    storage_location = input("[1] Generator Files Folder Path: ").strip()
    if not os.path.exists(storage_location):
        print("[!] That path does not exist.")
        return

    template_location = input("[2] Path to Folder of .docx Templates (not in [1]): ").strip()
    if not os.path.exists(template_location):
        print("[!] That path does not exist.")
        return

    templates_dir = os.path.join(storage_location, "templates")
    output_docx_dir = os.path.join(storage_location, "docxStorage")
    output_pdf_dir = os.path.join(storage_location, "outputPDFs")

    os.makedirs(templates_dir, exist_ok=True)
    os.makedirs(output_docx_dir, exist_ok=True)
    os.makedirs(output_pdf_dir, exist_ok=True)


    template_files = [f for f in os.listdir(template_location) if f.lower().endswith('.docx')]
    if not template_files:
        print("[!] No .docx files found in the template folder.")
        return

    print(f"Copying {len(template_files)} template files to {templates_dir}...")

    for file in template_files:
        src_path = os.path.join(template_location, file)
        dst_path = os.path.join(templates_dir, file)
        shutil.copy2(src_path, dst_path)

    print("[✓] Copy complete.")


    template_map = {}
    for i, filename in enumerate(template_files, start=1):
        while True:
            prompt = f"[>] Provide keyword for template {i} ('{filename}'): "
            keyword = input(prompt).strip()
            if keyword == "":
                print("[!] A keyword is necessary. Please enter a valid keyword.")
            else:
                template_map[keyword] = filename
                break  

    if not template_map:
        print("[!] No keywords assigned. Configuration aborted.")
        return

    swap_words = ["{{COMPANY_NAME}}", "{{ROLE}}", "{{DATE}}"]

   
    save_config(storage_location, templates_dir, template_map, output_docx_dir, output_pdf_dir, swap_words)

    print(f"\n[✓] Setup complete. You can now generate cover letters.")
    print(f"[✓] Output DOCX files will be saved to: {output_docx_dir}")
    print(f"[✓] Output PDF files will be saved to: {output_pdf_dir}")

def update_templates(custom_template_source=None):
    config = load_config()

    storage_location = config["storageLocation"]
    assigned_templates_dir = config["templateLocation"]
    output_docx_dir = config["outputDocxLocation"]
    output_pdf_dir = config["outputPdfLocation"]
    existing_map = config.get("templateKeywords", {})
    swap_words = config.get("swapWords", ["{{COMPANY_NAME}}", "{{ROLE}}", "{{DATE}}"])

    if custom_template_source:
        source_template_dir = custom_template_source
        if not os.path.exists(source_template_dir):
            print(f"[!] Provided source template directory does not exist: {source_template_dir}")
            return
    else:
        source_template_dir = assigned_templates_dir

    print(f"[>] Scanning for new templates in: {source_template_dir}")
    source_files = [f for f in os.listdir(source_template_dir) if f.lower().endswith(".docx")]
    known_files = set(existing_map.values())
    new_files = [f for f in source_files if f not in known_files]

    if not new_files:
        print("[✓] No new templates found to add.")
    else:
        print(f"[+] Found {len(new_files)} new template(s) to add.")

        added_count = 0
        for i, filename in enumerate(new_files, start=1):
            src_path = os.path.join(source_template_dir, filename)
            dst_path = os.path.join(assigned_templates_dir, filename)

            while True:
                keyword = input(f"[>] Provide keyword for new template {i} ('{filename}'): ").strip()

                if not keyword:
                    print("[!] Keyword cannot be empty. Please enter a valid keyword.")
                    continue

                if keyword in existing_map:
                    print(f"[!] Keyword already exists: '{keyword}' (associated with file '{existing_map[keyword]}'). Try a different one.")
                    continue

                if os.path.abspath(src_path) != os.path.abspath(dst_path):
                    shutil.copy2(src_path, dst_path)

                existing_map[keyword] = filename
                added_count += 1
                break

        if added_count:
            print(f"[✓] Added {added_count} new template(s).")
        else:
            print("[!] No templates were added.")

    print("\n[?] Do you want to add more {{PLACEHOLDERS}} ? [Y]es/[N]o (Default=No)")
    response = input("> ").strip().lower()

    while response in ("y", "yes"):
        while True:
            label = input("[>] Enter a LABEL for a new placeholder {{LABEL}}: ").strip().upper()
            if not label:
                print("[!] Placeholder label cannot be empty. Please enter a valid label.")
                continue

            placeholder = f"{{{{{label}}}}}"
            if placeholder in swap_words:
                print(f"[!] Placeholder {placeholder} already exists.")
            else:
                swap_words.append(placeholder)
                print(f"[✓] Added placeholder {placeholder}")
            break

        print("\n[?] Continue adding placeholders? [Y]es/[N]o (Default=No)")
        response = input("> ").strip().lower()

    save_config(storage_location, assigned_templates_dir, existing_map, output_docx_dir, output_pdf_dir, swap_words)




def show_config_summary():
    config = load_config()
    swap_words = config.get("swapWords", [])
    template_keywords = config.get("templateKeywords", {})
    storage_location = config.get("storageLocation", None)

    print("\n[✓] Current Configuration Summary:")

    print("\nStorage Location (Source Directory):")
    if storage_location:
        print(f"  {storage_location}")
    else:
        print("  [!] No storage location configured.")

    print("\nPlaceholders (Labels):")
    if swap_words:
        for word in swap_words:
            print(f"  - {word}")
    else:
        print("  [!] No placeholders found.")

    print("\nTemplate Keywords:")
    if template_keywords:
        for keyword, filename in template_keywords.items():
            print(f"  - '{keyword}' → {filename}")
    else:
        print("  [!] No templates configured.")


def clear_template_keywords():
    if not CONFIG_FILE.exists():
        print("[!] Config file not found, run HireMe --configure.")
        return

    with open(CONFIG_FILE, "r") as f:
        config = json.load(f)

    template_location = config.get("templateLocation")
    templates_dir = config.get("storageLocation")

    if not template_location or not templates_dir:
        print("[!] Invalid template or storage path in config.")
        return

    template_files = [f for f in os.listdir(template_location) if f.lower().endswith('.docx')]
    if not template_files:
        print("[!] No .docx files found in the template folder.")
        return


    config["templateKeywords"] = {}
    print("[✓] Cleared all existing template keywords.")

    template_map = {}
    for i, filename in enumerate(template_files, start=1):
        while True:
            prompt = f"[>] Provide keyword for template {i} ('{filename}'): "
            keyword = input(prompt).strip()
            if not keyword:
                print("[!] A keyword is necessary. Please enter a valid keyword.")
            elif keyword in template_map:
                print("[!] Duplicate keyword. Please enter a unique one.")
            else:
                template_map[keyword] = filename
                break

    if not template_map:
        print("[!] No keywords assigned. Configuration aborted.")
        return

    config["templateKeywords"] = template_map
    with open(CONFIG_FILE, "w") as f:
        json.dump(config, f, indent=4)

    print("[✓] Template keywords reassigned and saved.")




def reset_swap_words():
    config = load_config()
    current_placeholders = config.get("swapWords", ["{{COMPANY_NAME}}", "{{ROLE}}", "{{DATE}}"])
    print(f"Current placeholders: {current_placeholders}")

    choice = input("Reset to Default [Y]es / [C]lear All / [N]o (Default=No): ").strip().lower()

    if choice in ("y", "yes"):
        config["swapWords"] = ["{{COMPANY_NAME}}", "{{ROLE}}", "{{DATE}}"]
        print("[✓] Placeholders reset to default.")
    elif choice in ("c", "clearall","clear"):
        new_placeholders = []
        print("[>] Enter new placeholders one by one (without curly braces). Leave blank to finish.")
        while True:
            label = input("[>] Enter a LABEL for a new placeholder {{LABEL}}: ").strip().upper()
            if label == "":
                break
            placeholder = f"{{{{{label}}}}}"
            if placeholder in new_placeholders:
                print(f"[!] Placeholder {placeholder} already exists, skipped.")
            else:
                new_placeholders.append(placeholder)
                print(f"[✓] Added placeholder {placeholder}")

        config["swapWords"] = new_placeholders
        print(f"[✓] Placeholders updated: {config['swapWords']}")
    else:
        print("No changes made to placeholders.")
        return

    with open(CONFIG_FILE, "w") as f:
        json.dump(config, f, indent=4)


def generatorCore():
    config = load_config()
    storage_location = config["storageLocation"]
    templates_dir = config["templateLocation"]
    output_docx_dir = config["outputDocxLocation"]
    output_pdf_dir = config["outputPdfLocation"]


def clear_storage(storage_location):
    if not os.path.exists(storage_location):
        print(f"[!] Storage location '{storage_location}' does not exist.")
        return

    for filename in os.listdir(storage_location):
        file_path = os.path.join(storage_location, filename)
        try:
            if os.path.isfile(file_path) or os.path.islink(file_path):
                os.unlink(file_path)
            elif os.path.isdir(file_path):
                shutil.rmtree(file_path)
        except Exception as e:
            print(f"[!] Failed to delete {file_path}. Reason: {e}")

    if CONFIG_FILE.exists():
        try:
            CONFIG_FILE.unlink()
            print(f"[✓] Deleted configuration file at {CONFIG_FILE}")
        except Exception as e:
            print(f"[!] Failed to delete config file {CONFIG_FILE}. Reason: {e}")

    print(f"[✓] Cleared all contents inside {storage_location}")

def validate_config(config):

    if not config:
        return False

    source_location = config.get("storageLocation")
    templates_map = config.get("templateKeywords", {})
    swap_words = config.get("swapWords", [])

    if not source_location or not os.path.isdir(source_location):
        return False

    output_pdf_dir = os.path.join(source_location, "output.pdf")
    output_docx_dir = os.path.join(source_location, "output.docx")

    if not os.path.isdir(output_pdf_dir) or not os.path.isdir(output_docx_dir):
        return False

    docx_files = [f for f in os.listdir(output_docx_dir) if os.path.isfile(os.path.join(output_docx_dir, f))]
    if len(docx_files) == 0:
        return False

    if len(templates_map) != len(docx_files):
        return False

    if not swap_words or len(swap_words) < 1:
        return False

    return True


class CustomArgumentParser(argparse.ArgumentParser):
    def __init__(self, *args, **kwargs):
        if 'formatter_class' not in kwargs:
            kwargs['formatter_class'] = WideHelpFormatter
        super().__init__(*args, **kwargs)

    def error(self, message):
        args = sys.argv[1:]

        def is_config_valid():
            if not CONFIG_FILE.exists():
                return False
            try:
                with open(CONFIG_FILE, 'r') as f:
                    config = json.load(f)
            except Exception:
                return False

            required_paths = [
                "storageLocation",
                "templateLocation",
                "outputDocxLocation",
                "outputPdfLocation"
            ]
            for key in required_paths:
                val = config.get(key)
                if not val or not isinstance(val, str) or val.strip() == "":
                    return False
                if not os.path.exists(val):
                    return False

            template_keywords = config.get("templateKeywords")
            if not template_keywords or not isinstance(template_keywords, dict) or len(template_keywords) == 0:
                return False

            return True

        if '-G' in args or '--generate' in args:
            print(f"[!] {message}", file=sys.stderr)
            print("[>] usage: HireMe -G --company COMPANY_NAME --role ROLE_TITLE [--date DATE] [-T TEMPLATE_KEYWORD]", file=sys.stderr)
            sys.exit(2)

        elif not is_config_valid():
            print("[!] No configuration found", file=sys.stderr)
            print("[>] use: HireMe --configure", file=sys.stderr)
            sys.exit(2)

        else:
            super().error(message)
            print("[>] HireMe --help,-h ", file=sys.stderr)

class WideHelpFormatter(argparse.HelpFormatter):
    def __init__(self, prog):
        super().__init__(prog, width=120)
        self._max_help_position = 40 
        self._indent_increment = 2


####################################################################################################
####################################### MEAT AND POTATOES ##########################################
####################################################################################################
def generate_cover_letter(company, role, date, template_keyword):
    config = load_config()

    swap_words = config.get("swapWords", ["{{COMPANY_NAME}}", "{{ROLE}}", "{{DATE}}"])
    template_keywords = config.get("templateKeywords", {})
    template_dir = config.get("templateLocation")
    output_docx_dir = config.get("outputDocxLocation")
    output_pdf_dir = config.get("outputPdfLocation")

    if not template_dir or not os.path.exists(template_dir):
        print(f"[!] Template directory does not exist or is not configured: {template_dir}")
        return

    if not output_docx_dir or not os.path.exists(output_docx_dir):
        print(f"[!] Output DOCX directory does not exist or is not configured: {output_docx_dir}")
        return

    if not output_pdf_dir or not os.path.exists(output_pdf_dir):
        print(f"[!] Output PDF directory does not exist or is not configured: {output_pdf_dir}")
        return

    print(f"[>] Generating cover letter...")
    print(f"[>] Company: {company}")
    print(f"[>] Role: {role}")
    print(f"[>] Date: {date}")
    print(f"[>] Template Directory: {template_dir}")
    print(f"[>] Output DOCX Directory: {output_docx_dir}")
    print(f"[>] Output PDF Directory: {output_pdf_dir}")

    template_filename = template_keywords.get(template_keyword)
    if not template_filename:
        print(f"[!] No template found for keyword: '{template_keyword}'. Please update template keywords.")
        return

    template_path = os.path.join(template_dir, template_filename)

    if not os.path.exists(template_path):
        print(f"[!] Template file not found: {template_path}")
        return

    doc = Document(template_path)

    replacements = {
        "{{COMPANY_NAME}}": company,
        "{{ROLE}}": role,
        "{{DATE}}": date
    }
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for placeholder in swap_words:
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, replacements.get(placeholder, ""))

    safe_company = company.replace(" ", "").replace(".", "")
    safe_role = role.replace(" ", "").replace(".", "")
    docx_filename = f"CL_{safe_company}_{safe_role}.docx"
    pdf_filename = f"CL_{safe_company}_{safe_role}.pdf"

    edited_docx_path = os.path.join(output_docx_dir, docx_filename)
    output_pdf_path = os.path.join(output_pdf_dir, pdf_filename)

    doc.save(edited_docx_path)
    print(f"[✓] DOCX saved: {edited_docx_path}")

    convert(edited_docx_path)

    generated_pdf_path = edited_docx_path.replace(".docx", ".pdf")

    timeout = 10
    start_time = time.time()
    while not os.path.exists(generated_pdf_path):
        if time.time() - start_time > timeout:
            print(f"[!] PDF was not generated at: {generated_pdf_path}")
            break
        time.sleep(0.5)
    else:
        shutil.move(generated_pdf_path, output_pdf_path)
        print(f"[✓] PDF saved: {output_pdf_path}")
####################################################################################################
####################################################################################################
####################################################################################################


def main():

    parser = CustomArgumentParser(description="HireMe Cover Letter Generator")

    parser.add_argument("--show", action="store_true", help="show current template keywords and placeholders")
    parser.add_argument("--configure", action="store_true", help="configure HireMe with storage and template locations")
    parser.add_argument("-G", "--generate", action="store_true", help="generate cover letter") 
    parser.add_argument("--company", help="company Name")
    parser.add_argument("--role", help="role title")
    parser.add_argument("--date", help="date [OPTIONAL, DEFAULT is today's date]")
    parser.add_argument("-T", "--template", help="template keyword to use when generating (required if multiple templates exist)")
    parser.add_argument("--clean", action="store_true", help="clear the storage directory contents")
    parser.add_argument("--update", action="store_true", help="scan for and import new templates")
    parser.add_argument("-S", "--source", help="optional custom path for template update (used with --update)")
    parser.add_argument("-L", "--resetlabels", action="store_true", help="reset or clear the placeholder labels")
    parser.add_argument("-K", "--resetkwds", action="store_true", help="clear all template keywords in configuration")
    parser.add_argument("--reset", action="store_true", help="reset everything: clean, reconfigure, and reset placeholders")
    parser.add_argument('--info', action='store_true', help='objective truth')

    

    args = parser.parse_args()

    if len(sys.argv) == 1:
        print("[>] use: HireMe --help")
        return

    if args.reset:
        other_flags_used = any([
            args.generate, args.configure, args.clean, args.update,
            args.company, args.role, args.date,
            args.source, args.resetlabels, args.resetkwds
        ])
        if other_flags_used:
            print("[!] --reset must be used alone with no other flags.")
            exit(1)

        print("[!] This will fully reset and reconfigure HireMe")
        confirm = input("[>] Continue? [Y]es/[N]o (Default=No): ").strip().lower()
        if confirm not in ("y", "yes"):
            print("[✓] Aborted.")
            return

        config = load_config()
        storage_location = config.get("storageLocation")
        clear_storage(storage_location)
        configureLess()
        reset_swap_words()
        return

    if args.clean or args.configure:
        other_flags = any([
            args.generate, args.company, args.role, args.date,
            args.update, args.source, args.resetlabels, args.resetkwds, args.reset
        ])
        if other_flags:
            print("[!] --clean and --configure must be used alone with no other flags.")
            exit(1)

    if args.clean:
        config = load_config() 
        storage_location = config.get("storageLocation")

        print("[!] This will delete all contents within the Source Directory and Configuration File.")
        confirm = input("[>] Continue? [Y]es/[N]o (Default=No): ").strip().lower()

        if confirm not in ("y", "yes"):
            print("[✓] Aborted. No files were deleted.")
            return

        clear_storage(storage_location)

    elif args.configure:
        configure()  

    
    if args.info:
        other_flags_used = any([
            args.configure, args.generate, args.clean, args.update, args.reset,
            args.resetlabels, args.resetkwds, args.company, args.role, args.date, args.template, args.source
        ])

        if other_flags_used:
            print("[!] The --info flag must be used by itself.")
            sys.exit(1)

        print("\nPhysics is better than programming (FORTRAN >> C++)\n")
        return
  
    if args.update or args.source:
        other_flags = any([
            args.generate, args.company, args.role, args.date,
            args.clean, args.configure, args.resetlabels, args.resetkwds, args.reset
        ])
        if other_flags:
            print("[!] --update (and --source) must be used alone.")
            exit(1)

        update_templates(custom_template_source=args.source)
        return

    if args.resetlabels or args.resetkwds:
        other_flags = any([
            args.generate, args.company, args.role, args.date,
            args.clean, args.configure, args.update, args.source, args.reset
        ])
        if other_flags:
            print("[!] --resetlabels (-L) and --resetkwds (-K) can only be used together, with no other flags.")
            exit(1)

        reset_parts = []
        if args.resetkwds:
            reset_parts.append("keywords")
        if args.resetlabels:
            reset_parts.append("{{LABELS}}")
        
        print(f"[!] This will reset the configured ({', '.join(reset_parts)}).")
        confirm = input("[>] Continue? [Y]es/[N]o (Default=No): ").strip().lower()
        if confirm not in ("y", "yes"):
            print("[x] Operation aborted.")
            return

        if args.resetkwds:
            clear_template_keywords()
        if args.resetlabels:
            reset_swap_words()
        return



    if args.show:
        other_flags = any([
            args.generate, args.configure, args.clean, args.update, args.source,
            args.company, args.role, args.date, args.resetlabels, args.resetkwds, args.reset, args.template
        ])
        if other_flags:
            print("[!] --show must be used alone.")
            exit(1)

        show_config_summary()
        return

    if args.generate:
        if not args.company or not args.role:
            print("[!] --company and --role are required when using -G/--generate.")
            print("[>] usage: HireMe -G --company 'COMPANY_NAME' --role 'ROLE_TITLE' [--date 'DATE'] [-T 'TEMPLATE_KEYWORD']")
            exit(1)

        config = load_config()
        storage_location = config.get("storageLocation")
        template_keywords = config.get("templateKeywords", {})

        if not storage_location or not os.path.exists(storage_location):
            print(f"[!] Configured storage location '{storage_location}' does not exist. Run --configure.")
            exit(1)

        if len(template_keywords) > 1 and not args.template:
            print("[!] Multiple templates found. You must specify one using --template.")
            print(f"[>] Available templates: {', '.join(template_keywords.keys())}")
            exit(1)


        selected_template = args.template
        if len(template_keywords) > 1:
            if selected_template not in template_keywords:
                print(f"[!] Invalid template keyword: '{selected_template}'")
                print(f"[>] Available templates: {', '.join(template_keywords.keys())}")
                exit(1)
        else:
            selected_template = list(template_keywords.keys())[0] if template_keywords else None

        generate_cover_letter(
            args.company,
            args.role,
            args.date or datetime.now().strftime('%m/%d/%y'),
            selected_template
        )
        return

    else:
        print()




if __name__ == "__main__":
    main()